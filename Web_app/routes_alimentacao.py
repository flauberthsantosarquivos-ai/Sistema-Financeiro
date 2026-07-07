from __future__ import annotations

import csv
import re
import shutil
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi import APIRouter, File, Form, Query, Request, UploadFile
from fastapi.responses import HTMLResponse
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook

from core.financeiro.alimentacao_google import (
    CONTAS_ALIMENTACAO,
    identificar_conta_alimentacao_por_dados_bancarios,
    importar_movimentacoes_extrato,
    montar_resumo_alimentacao,
    ler_contas,
)
from core.financeiro.configuracao_sistema import obter_configuracao_sistema
from core.financeiro.leitor_extrato import carregar_previa_extrato, salvar_previa_extrato
from core.financeiro.patrimonio_google import atualizar_saldo_pagbank_por_alimentacao


router = APIRouter()
templates = Jinja2Templates(directory="Web_app/templates")
PASTA_UPLOAD_ALIMENTACAO = Path("uploads/alimentacao")
PASTA_UPLOAD_ALIMENTACAO.mkdir(parents=True, exist_ok=True)


# Leitura exclusiva do módulo Alimentação/PagBank.
# Não altera o leitor_extrato.py usado pelo módulo geral de Extratos.
COLUNAS_PAGBANK = {
    "data": {"data"},
    "descricao": {"descricao", "descrição", "historico", "histórico"},
    "entradas": {"entradas", "entrada", "creditos", "créditos", "credito", "crédito"},
    "saidas": {"saidas", "saídas", "saida", "saída", "debitos", "débitos", "debito", "débito"},
    "saldo": {"saldo", "saldo atual", "saldo final"},
}


def obter_referencia(ano: str | None, mes: str | None) -> tuple[str, str]:
    hoje = date.today()
    return str(ano or hoje.year), str(mes or hoje.month).zfill(2)


def obter_referencia_do_extrato(movimentacoes: list[dict]) -> tuple[str, str]:
    """Obtém ano e mês pela data mais recente do extrato."""
    datas: list[datetime] = []
    for mov in movimentacoes:
        texto_data = str(mov.get("data") or mov.get("DATA") or "").strip()
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y"):
            try:
                datas.append(datetime.strptime(texto_data[:10], formato))
                break
            except ValueError:
                continue
    if not datas:
        return obter_referencia(None, None)
    referencia = max(datas)
    return str(referencia.year), f"{referencia.month:02d}"


def _normalizar(valor: Any) -> str:
    texto = str(valor or "").strip().lower()
    texto = unicodedata.normalize("NFKD", texto).encode("ASCII", "ignore").decode("ASCII")
    texto = re.sub(r"\s+", " ", texto)
    return texto


def _converter_valor(valor: Any) -> float:
    if valor is None or valor == "":
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if not texto:
        return 0.0
    negativo_parenteses = texto.startswith("(") and texto.endswith(")")
    texto = texto.replace("(", "").replace(")", "")

    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")

    texto = re.sub(r"[^0-9.-]", "", texto)
    try:
        numero = float(texto)
    except ValueError:
        return 0.0
    return -abs(numero) if negativo_parenteses else numero


def _formatar_valor_br(valor: float) -> str:
    return f"{abs(valor):.2f}".replace(".", ",")


def _formatar_data(valor: Any) -> str:
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")

    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto[:10], formato).strftime("%d/%m/%Y")
        except ValueError:
            pass
    return texto


def _texto_do_cabecalho_excel(caminho: Path) -> str:
    livro = load_workbook(caminho, read_only=True, data_only=True)
    planilha = livro[livro.sheetnames[0]]
    linhas = []
    for linha in planilha.iter_rows(min_row=1, max_row=35, values_only=True):
        texto_linha = " ".join(str(c).strip() for c in linha if c is not None and str(c).strip())
        if texto_linha:
            linhas.append(texto_linha)
    livro.close()
    return "\n".join(linhas)


def _texto_do_cabecalho_csv(caminho: Path) -> str:
    for encoding in ("utf-8-sig", "latin-1", "cp1252"):
        try:
            return caminho.read_text(encoding=encoding, errors="ignore")[:8000]
        except UnicodeDecodeError:
            continue
    return caminho.read_text(encoding="utf-8", errors="ignore")[:8000]


def _extrair_valor_cabecalho(texto: str, rotulo: str) -> str:
    correspondencia = re.search(rf"{rotulo}\s*:?\s*([^\r\n]+)", texto, flags=re.IGNORECASE)
    if not correspondencia:
        return ""
    valor = correspondencia.group(1).strip()
    return re.split(r"\s+(?:Banco|Ag[eê]ncia|Conta|Per[ií]odo)\s*: ?", valor, maxsplit=1, flags=re.IGNORECASE)[0].strip()


def identificar_conta_no_arquivo(caminho: Path) -> dict:
    extensao = caminho.suffix.lower()
    if extensao in {".xlsx", ".xls"}:
        texto = _texto_do_cabecalho_excel(caminho)
    elif extensao == ".csv":
        texto = _texto_do_cabecalho_csv(caminho)
    else:
        raise ValueError("Formato não suportado para identificação da conta.")

    banco = _extrair_valor_cabecalho(texto, r"Banco")
    agencia = _extrair_valor_cabecalho(texto, r"Ag[eê]ncia")
    conta = _extrair_valor_cabecalho(texto, r"Conta")
    conta_interna = identificar_conta_alimentacao_por_dados_bancarios(banco, agencia, conta)
    return {"conta": conta_interna, "banco": banco, "agencia": agencia, "numero_conta": conta}


def _localizar_cabecalho_pagbank(linhas: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    for indice, linha in enumerate(linhas[:60]):
        mapa = {}
        for coluna, celula in enumerate(linha):
            cabecalho = _normalizar(celula)
            for chave, alternativas in COLUNAS_PAGBANK.items():
                if cabecalho in alternativas:
                    mapa[chave] = coluna
        if {"data", "descricao", "entradas", "saidas"}.issubset(mapa):
            return indice, mapa
    raise ValueError("Não encontrei a tabela PagBank. Esperado: Data, Descrição, Entradas e Saídas.")


def _ler_extrato_pagbank_excel(caminho: Path) -> dict:
    livro = load_workbook(caminho, data_only=True)
    planilha = livro[livro.sheetnames[0]]
    linhas = list(planilha.iter_rows(values_only=True))
    livro.close()

    indice_cabecalho, colunas = _localizar_cabecalho_pagbank(linhas)
    movimentos: list[dict] = []
    saldo_final: float | None = None
    data_saldo = ""

    for numero_linha, linha in enumerate(linhas[indice_cabecalho + 1 :], start=indice_cabecalho + 2):
        data = _formatar_data(linha[colunas["data"]] if colunas["data"] < len(linha) else "")
        descricao = str(
            linha[colunas["descricao"]]
            if colunas["descricao"] < len(linha) and linha[colunas["descricao"]] is not None
            else ""
        ).strip()
        entrada = _converter_valor(linha[colunas["entradas"]] if colunas["entradas"] < len(linha) else "")
        saida = _converter_valor(linha[colunas["saidas"]] if colunas["saidas"] < len(linha) else "")

        # O saldo final pode aparecer na última movimentação ou em uma linha "Saldo do dia".
        if "saldo" in colunas and colunas["saldo"] < len(linha):
            saldo_bruto = linha[colunas["saldo"]]
            if data and saldo_bruto not in (None, ""):
                saldo_final = abs(_converter_valor(saldo_bruto))
                data_saldo = data

        if not data and not descricao and entrada == 0 and saida == 0:
            continue

        descricao_normalizada = _normalizar(descricao)
        if "saldo do dia" in descricao_normalizada or descricao_normalizada.startswith("saldo"):
            continue
        if not data or not descricao:
            continue

        if abs(entrada) > 0:
            valor = abs(entrada)
            tipo = "RECEITA"
        elif abs(saida) > 0:
            valor = -abs(saida)
            tipo = "DESPESA"
        else:
            continue

        movimentos.append({
            "indice": len(movimentos),
            "linha_original": numero_linha,
            "data": data,
            "descricao": descricao,
            "valor": valor,
            "valor_fmt": _formatar_valor_br(valor),
            "tipo": tipo,
            "categoria": "ALIMENTAÇÃO",
            "subcategoria": "PAGBANK",
            "situacao": "RECEBIDO" if tipo == "RECEITA" else "PAGO",
            "forma_pagamento": "CONTA",
            "observacao": f"Importado do extrato PagBank: {caminho.name}",
        })

    return {
        "movimentacoes": movimentos,
        "saldo_final": saldo_final,
        "data_saldo": data_saldo,
    }

def _ler_extrato_pagbank_csv(caminho: Path) -> dict:
    for encoding in ("utf-8-sig", "latin-1", "cp1252"):
        try:
            with caminho.open("r", encoding=encoding, newline="") as arquivo:
                amostra = arquivo.read(4096)
                arquivo.seek(0)
                delimitador = ";" if amostra.count(";") >= amostra.count(",") else ","
                leitor = csv.reader(arquivo, delimiter=delimitador)
                linhas = [tuple(linha) for linha in leitor]
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Não foi possível ler o CSV do PagBank.")

    indice_cabecalho, colunas = _localizar_cabecalho_pagbank(linhas)
    movimentos: list[dict] = []
    saldo_final: float | None = None
    data_saldo = ""

    for numero_linha, linha in enumerate(linhas[indice_cabecalho + 1 :], start=indice_cabecalho + 2):
        data = _formatar_data(linha[colunas["data"]] if colunas["data"] < len(linha) else "")
        descricao = str(linha[colunas["descricao"]] if colunas["descricao"] < len(linha) else "").strip()
        entrada = _converter_valor(linha[colunas["entradas"]] if colunas["entradas"] < len(linha) else "")
        saida = _converter_valor(linha[colunas["saidas"]] if colunas["saidas"] < len(linha) else "")

        if "saldo" in colunas and colunas["saldo"] < len(linha):
            saldo_bruto = linha[colunas["saldo"]]
            if data and saldo_bruto not in (None, ""):
                saldo_final = abs(_converter_valor(saldo_bruto))
                data_saldo = data

        if not data or not descricao or "saldo do dia" in _normalizar(descricao):
            continue
        if abs(entrada) > 0:
            valor, tipo = abs(entrada), "RECEITA"
        elif abs(saida) > 0:
            valor, tipo = -abs(saida), "DESPESA"
        else:
            continue

        movimentos.append({
            "indice": len(movimentos),
            "linha_original": numero_linha,
            "data": data,
            "descricao": descricao,
            "valor": valor,
            "valor_fmt": _formatar_valor_br(valor),
            "tipo": tipo,
            "categoria": "ALIMENTAÇÃO",
            "subcategoria": "PAGBANK",
            "situacao": "RECEBIDO" if tipo == "RECEITA" else "PAGO",
            "forma_pagamento": "CONTA",
            "observacao": f"Importado do extrato PagBank: {caminho.name}",
        })

    return {
        "movimentacoes": movimentos,
        "saldo_final": saldo_final,
        "data_saldo": data_saldo,
    }

def processar_extrato_alimentacao_pagbank(caminho: Path) -> dict:
    extensao = caminho.suffix.lower()
    if extensao in {".xlsx", ".xls"}:
        resultado = _ler_extrato_pagbank_excel(caminho)
    elif extensao == ".csv":
        resultado = _ler_extrato_pagbank_csv(caminho)
    else:
        raise ValueError("Formato não suportado. Envie CSV, XLS ou XLSX.")

    resultado["total_linhas"] = len(resultado["movimentacoes"])
    return resultado



def _mes_sigla(numero_mes: str) -> str:
    meses = {
        "01": "JAN", "02": "FEV", "03": "MAR", "04": "ABR",
        "05": "MAI", "06": "JUN", "07": "JUL", "08": "AGO",
        "09": "SET", "10": "OUT", "11": "NOV", "12": "DEZ",
    }
    return meses.get(str(numero_mes or "").zfill(2), "")



def resolver_periodo_alimentacao(
    ano: str,
    mes: str,
    periodo: str | None,
    data_inicio: str | None,
    data_fim: str | None,
) -> tuple[str, str, str]:
    """Converte a escolha de semana/período em um intervalo ISO dentro do mês."""
    periodo_final = str(periodo or "mes").strip().lower()
    try:
        ano_num = int(ano)
        mes_num = int(mes)
    except (TypeError, ValueError):
        return "", "", "Mês inteiro"

    ultimo_dia = 31
    while True:
        try:
            date(ano_num, mes_num, ultimo_dia)
            break
        except ValueError:
            ultimo_dia -= 1

    faixas = {
        "semana_1": (1, 7, "1ª semana · dias 1 a 7"),
        "semana_2": (8, 14, "2ª semana · dias 8 a 14"),
        "semana_3": (15, 21, "3ª semana · dias 15 a 21"),
        "semana_4": (22, ultimo_dia, f"4ª semana · dias 22 a {ultimo_dia}"),
    }

    if periodo_final in faixas:
        inicio_dia, fim_dia, descricao = faixas[periodo_final]
        return (
            date(ano_num, mes_num, inicio_dia).isoformat(),
            date(ano_num, mes_num, fim_dia).isoformat(),
            descricao,
        )

    if periodo_final == "personalizado":
        inicio = str(data_inicio or "").strip()
        fim = str(data_fim or "").strip()
        return inicio, fim, "Período personalizado"

    return "", "", "Mês inteiro"

def montar_contexto(request: Request, ano: str, mes: str, *, etapa: str = "painel", mensagem: str | None = None, erro: str | None = None, resumo: dict | None = None, resultado: dict | None = None, temp_id: str = "", conta_selecionada: str = "", periodo: str = "mes", data_inicio: str = "", data_fim: str = "", periodo_descricao: str = "Mês inteiro", conta_filtro: str = ""):
    if resumo is None:
        resumo = montar_resumo_alimentacao(ano, mes, data_inicio=data_inicio, data_fim=data_fim, conta_filtro=conta_filtro)
    return {"config": obter_configuracao_sistema(), "ano": ano, "mes": mes, "etapa": etapa, "mensagem": mensagem, "erro": erro, "resumo": resumo, "resultado": resultado or {}, "temp_id": temp_id, "contas_alimentacao": CONTAS_ALIMENTACAO, "conta_selecionada": conta_selecionada, "periodo": periodo, "data_inicio": data_inicio, "data_fim": data_fim, "periodo_descricao": periodo_descricao, "conta_filtro": conta_filtro}


@router.get("/financeiro/alimentacao", response_class=HTMLResponse)
async def alimentacao_get(
    request: Request,
    ano: str | None = Query(None),
    mes: str | None = Query(None),
    salvo: str | None = Query(None),
    periodo: str | None = Query(None),
    data_inicio: str | None = Query(None),
    data_fim: str | None = Query(None),
    conta: str | None = Query(None),
):
    ano_final, mes_final = obter_referencia(ano, mes)
    inicio_final, fim_final, periodo_descricao = resolver_periodo_alimentacao(
        ano_final, mes_final, periodo, data_inicio, data_fim
    )
    conta_final = str(conta or "").strip().upper()
    if conta_final not in CONTAS_ALIMENTACAO:
        conta_final = ""
    mensagem = "" if not salvo else "Extrato de alimentação importado com sucesso."
    try:
        contexto = montar_contexto(
            request,
            ano_final,
            mes_final,
            mensagem=mensagem,
            periodo=str(periodo or "mes"),
            data_inicio=inicio_final,
            data_fim=fim_final,
            periodo_descricao=periodo_descricao,
            conta_filtro=conta_final,
        )
    except Exception as e:
        contexto = montar_contexto(
            request,
            ano_final,
            mes_final,
            erro=f"Erro ao carregar alimentação: {e}",
            resumo={"contas": [], "movimentacoes": [], "total_saldo_fmt": "R$ 0,00", "total_entradas_fmt": "R$ 0,00", "total_saidas_fmt": "R$ 0,00", "qtd_movimentacoes": 0, "analise": {"qtd_compras": 0, "ticket_medio_fmt": "R$ 0,00", "dias_com_gasto": 0, "gasto_por_conta": [], "maiores_gastos": [], "evolucao_diaria": []}},
            periodo=str(periodo or "mes"),
            data_inicio=inicio_final,
            data_fim=fim_final,
            periodo_descricao=periodo_descricao,
            conta_filtro=conta_final,
        )
    return templates.TemplateResponse(request=request, name="alimentacao.html", context=contexto)


@router.post("/financeiro/alimentacao/sincronizar-patrimonio", response_class=HTMLResponse)
async def alimentacao_sincronizar_patrimonio(request: Request):
    ano_final, mes_final = obter_referencia(None, None)
    try:
        contas = ler_contas()
        mes_patrimonio = _mes_sigla(mes_final)
        if not mes_patrimonio:
            raise ValueError("Mês de referência inválido para sincronização patrimonial.")

        atualizadas = []
        for item in contas:
            conta = str(item.get("conta", "")).strip().upper()
            saldo = item.get("saldo_atual", 0)
            data_saldo = str(item.get("data_saldo", "") or "").strip()
            resultado = atualizar_saldo_pagbank_por_alimentacao(
                ano=ano_final,
                mes=mes_patrimonio,
                subdivisao=conta,
                saldo_final=saldo,
                data_fim=data_saldo,
            )
            atualizadas.append(f"{conta}: {resultado['saldo_final_fmt']}")

        resumo = montar_resumo_alimentacao(ano_final, mes_final)
        mensagem = "Saldos sincronizados com o Patrimônio: " + " | ".join(atualizadas)
        return templates.TemplateResponse(
            request=request,
            name="alimentacao.html",
            context=montar_contexto(request, ano_final, mes_final, mensagem=mensagem, resumo=resumo),
        )
    except Exception as e:
        return templates.TemplateResponse(
            request=request,
            name="alimentacao.html",
            context=montar_contexto(
                request, ano_final, mes_final,
                erro=f"Erro ao sincronizar saldos com o Patrimônio: {e}",
            ),
        )


@router.post("/financeiro/alimentacao/previsualizar", response_class=HTMLResponse)
async def alimentacao_previsualizar(
    request: Request,
    arquivo: UploadFile = File(...),
    ano: str = Form(""),
    mes: str = Form(""),
    saldo_atual: str = Form(""),
):
    # A tela não envia Ano/Mês: a referência vem das datas lidas no extrato.
    ano_final, mes_final = obter_referencia(ano, mes)
    nome = arquivo.filename or ""
    if not nome.lower().endswith((".csv", ".xlsx", ".xls")):
        return templates.TemplateResponse(request=request, name="alimentacao.html", context=montar_contexto(request, ano_final, mes_final, erro="Formato não suportado. Envie CSV, XLS ou XLSX."))
    try:
        caminho = PASTA_UPLOAD_ALIMENTACAO / f"{uuid4().hex}_{nome}"
        with caminho.open("wb") as buffer:
            shutil.copyfileobj(arquivo.file, buffer)

        identificacao = identificar_conta_no_arquivo(caminho)
        conta_final = identificacao["conta"]
        if conta_final not in CONTAS_ALIMENTACAO:
            raise ValueError(f"Não foi possível reconhecer a conta pelo cabeçalho. Banco: {identificacao['banco'] or 'não encontrado'} | Agência: {identificacao['agencia'] or 'não encontrada'} | Conta: {identificacao['numero_conta'] or 'não encontrada'}")

        resultado = processar_extrato_alimentacao_pagbank(caminho)
        movimentos = [
            {**mov, "conta": conta_final, "importar": "SIM", "categoria_interna": conta_final}
            for mov in resultado["movimentacoes"]
        ]
        ano_final, mes_final = obter_referencia_do_extrato(movimentos)
        saldo_lido = resultado.get("saldo_final")
        saldo_para_salvar = saldo_atual if str(saldo_atual or "").strip() else saldo_lido
        previa = {
            "origem": "ALIMENTACAO",
            "ano": ano_final,
            "mes": mes_final,
            "conta": conta_final,
            "saldo_atual": saldo_para_salvar,
            "data_saldo": resultado.get("data_saldo") or "",
            "movimentacoes": movimentos,
            "nome_arquivo": nome,
            "identificacao_bancaria": identificacao,
        }
        temp_id = salvar_previa_extrato(previa)
        resumo = montar_resumo_alimentacao(ano_final, mes_final)
        mensagem = f"Extrato lido com sucesso. Conta identificada automaticamente: {conta_final}. {len(movimentos)} movimentação(ões) encontrada(s)."
        return templates.TemplateResponse(request=request, name="alimentacao.html", context=montar_contexto(request, ano_final, mes_final, etapa="previa", mensagem=mensagem, resumo=resumo, resultado=previa, temp_id=temp_id, conta_selecionada=conta_final))
    except Exception as e:
        return templates.TemplateResponse(request=request, name="alimentacao.html", context=montar_contexto(request, ano_final, mes_final, erro=f"Erro ao ler extrato: {e}"))


@router.post("/financeiro/alimentacao/importar", response_class=HTMLResponse)
async def alimentacao_importar(request: Request):
    form = await request.form()
    temp_id = str(form.get("temp_id", "")).strip()
    try:
        previa = carregar_previa_extrato(temp_id)
        if previa.get("origem") != "ALIMENTACAO":
            raise ValueError("Prévia inválida para o módulo Alimentação.")
        ano_final, mes_final = obter_referencia(previa.get("ano"), previa.get("mes"))
        conta = str(previa.get("conta", "")).strip().upper()
        if conta not in CONTAS_ALIMENTACAO:
            raise ValueError("A conta identificada na prévia não é válida.")
        selecionadas = []
        for mov in previa.get("movimentacoes", []):
            indice = str(mov.get("indice"))
            if form.get(f"importar_{indice}") == "SIM":
                selecionadas.append({**mov, "tipo": str(mov.get("tipo", "")).strip().upper(), "conta": conta, "categoria_interna": str(form.get(f"categoria_{indice}", conta)).strip() or conta})
        saldo_atual = previa.get("saldo_atual", "")
        data_saldo = str(previa.get("data_saldo", "") or "").strip()
        resultado_importacao = importar_movimentacoes_extrato(
            conta=conta,
            movimentacoes=selecionadas,
            saldo_atual=saldo_atual,
            data_saldo=data_saldo,
        )

        mensagem = (
            f"Importação concluída para {conta}: "
            f"{resultado_importacao['importadas']} movimentação(ões) nova(s). "
            f"{resultado_importacao['duplicadas']} duplicada(s) ignorada(s)."
        )

        # A sincronização é automática inclusive quando a importação for composta só de duplicadas.
        if saldo_atual not in (None, ""):
            mes_patrimonio = _mes_sigla(mes_final)
            if not mes_patrimonio:
                raise ValueError("Mês de referência inválido para atualização patrimonial.")
            resultado_patrimonio = atualizar_saldo_pagbank_por_alimentacao(
                ano=ano_final,
                mes=mes_patrimonio,
                subdivisao=conta,
                saldo_final=saldo_atual,
                data_fim=data_saldo,
            )
            mensagem += (
                f" Saldo sincronizado automaticamente no Patrimônio: "
                f"PAGBANK / {conta} = {resultado_patrimonio['saldo_final_fmt']}."
            )
        else:
            mensagem += " Não foi localizado saldo final no extrato; o Patrimônio não foi alterado."

        resumo = montar_resumo_alimentacao(ano_final, mes_final)
        return templates.TemplateResponse(request=request, name="alimentacao.html", context=montar_contexto(request, ano_final, mes_final, mensagem=mensagem, resumo=resumo))
    except Exception as e:
        ano_final, mes_final = obter_referencia(None, None)
        return templates.TemplateResponse(request=request, name="alimentacao.html", context=montar_contexto(request, ano_final, mes_final, erro=f"Erro ao importar extrato de alimentação: {e}"))
