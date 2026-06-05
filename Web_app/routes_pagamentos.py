"""
Rotas do Módulo de Pagamentos

Arquivo sugerido:
Web_app/routes_pagamentos.py

Integração no app.py:
from Web_app.routes_pagamentos import router as pagamentos_router
app.include_router(pagamentos_router)
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Form, Query, Request
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates

from core.financeiro.pagamentos_google import (
    MESES_ORDEM,
    STATUS_PAGAMENTO,
    conciliar_pagamentos_com_lancamentos,
    duplicar_pagamento,
    duplicar_pagamentos_mes,
    excluir_pagamento,
    listar_pagamentos,
    montar_resumo_pagamentos,
    obter_opcoes_categorias,
    salvar_pagamento,
)


router = APIRouter(prefix="/financeiro", tags=["Financeiro - Pagamentos"])
templates = Jinja2Templates(directory="Web_app/templates")


def redirect_pagamentos(params: Optional[str] = None):
    url = "/financeiro/pagamentos"
    if params:
        url += f"?{params}"
    return RedirectResponse(url=url, status_code=303)


def filtros_padrao(
    ano: Optional[int] = None,
    mes: Optional[str] = None,
    status: Optional[str] = None,
    tipo: Optional[str] = None,
    categoria: Optional[str] = None,
):
    hoje = date.today()

    return {
        "ano": str(ano or hoje.year),
        "mes": (mes or MESES_ORDEM[hoje.month - 1]).upper(),
        "status": (status or "").upper(),
        "tipo": (tipo or "").upper(),
        "categoria": (categoria or "").upper(),
    }



def montar_filtros_relatorio(
    ano: Optional[int] = None,
    mes: Optional[str] = None,
    status: Optional[str] = None,
    tipo: Optional[str] = None,
    categoria: Optional[str] = None,
) -> dict:
    return filtros_padrao(
        ano=ano,
        mes=mes,
        status=status,
        tipo=tipo,
        categoria=categoria,
    )


def nome_arquivo_relatorio(filtros: dict, extensao: str) -> str:
    ano = str(filtros.get("ano") or "todos")
    mes = str(filtros.get("mes") or "todos").lower()
    data_geracao = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"relatorio_pagamentos_{ano}_{mes}_{data_geracao}.{extensao}"


def texto_filtro(valor: str) -> str:
    texto = str(valor or "").strip()
    return texto if texto else "Todos"


@router.get("/pagamentos")
async def tela_pagamentos(
    request: Request,
    ano: Optional[int] = Query(None),
    mes: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    tipo: Optional[str] = Query(None),
    categoria: Optional[str] = Query(None),
    mensagem: Optional[str] = Query(None),
    erro: Optional[str] = Query(None),
):
    filtros = filtros_padrao(ano=ano, mes=mes, status=status, tipo=tipo, categoria=categoria)

    try:
        pagamentos = listar_pagamentos(filtros)
        resumo = montar_resumo_pagamentos(pagamentos)
        categorias = obter_opcoes_categorias()
    except Exception as exc:
        pagamentos = []
        resumo = montar_resumo_pagamentos([])
        categorias = {}
        erro = erro or f"Erro ao carregar pagamentos: {exc}"

    tipos_opcoes = [""] + sorted(categorias.keys())
    meses_opcoes = [
        ("JAN", "Janeiro"),
        ("FEV", "Fevereiro"),
        ("MAR", "Março"),
        ("ABR", "Abril"),
        ("MAI", "Maio"),
        ("JUN", "Junho"),
        ("JUL", "Julho"),
        ("AGO", "Agosto"),
        ("SET", "Setembro"),
        ("OUT", "Outubro"),
        ("NOV", "Novembro"),
        ("DEZ", "Dezembro"),
    ]

    return templates.TemplateResponse(
        request=request,
        name="pagamentos.html",
        context={
            "request": request,
            "filtros": filtros,
            "pagamentos": pagamentos,
            "resumo": resumo,
            "categorias": categorias,
            "status_opcoes": [""] + STATUS_PAGAMENTO,
            "tipos_opcoes": tipos_opcoes,
            "meses_opcoes": meses_opcoes,
            "mensagem": mensagem,
            "erro": erro,
        },
    )


@router.get("/pagamentos/relatorio/pdf")
async def relatorio_pagamentos_pdf(
    ano: Optional[int] = Query(None),
    mes: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    tipo: Optional[str] = Query(None),
    categoria: Optional[str] = Query(None),
):
    filtros = montar_filtros_relatorio(
        ano=ano,
        mes=mes,
        status=status,
        tipo=tipo,
        categoria=categoria,
    )

    pagamentos = listar_pagamentos(filtros)
    resumo = montar_resumo_pagamentos(pagamentos)

    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except Exception as exc:
        raise RuntimeError(
            "Para gerar PDF, instale a biblioteca reportlab: pip install reportlab"
        ) from exc

    def status_visual(item: dict) -> str:
        return str(item.get("STATUS_CALCULADO") or item.get("STATUS") or "").strip().upper()

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=0.75 * cm,
        leftMargin=0.75 * cm,
        topMargin=0.75 * cm,
        bottomMargin=0.75 * cm,
    )

    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle(
        "TituloRelatorioPagamentosBonito",
        parent=estilos["Title"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=6,
    )
    subtitulo = ParagraphStyle(
        "SubtituloRelatorioPagamentosBonito",
        parent=estilos["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica",
        fontSize=8.2,
        leading=10,
        textColor=colors.HexColor("#475569"),
    )
    celula = ParagraphStyle(
        "CelulaRelatorioPagamentosBonito",
        parent=estilos["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica",
        fontSize=7.4,
        leading=9.0,
        textColor=colors.HexColor("#111827"),
    )
    celula_negrito = ParagraphStyle(
        "CelulaNegritoRelatorioPagamentosBonito",
        parent=celula,
        fontName="Helvetica-Bold",
    )
    celula_pago = ParagraphStyle(
        "CelulaPagoRelatorioPagamentosBonito",
        parent=celula,
        fontName="Helvetica-Bold",
        textColor=colors.HexColor("#065f46"),
    )

    elementos = []
    elementos.append(Paragraph("Relatório de Pagamentos", titulo))
    elementos.append(
        Paragraph(
            f"Ano: <b>{texto_filtro(filtros.get('ano'))}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Mês: <b>{texto_filtro(filtros.get('mes'))}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Status: <b>{texto_filtro(filtros.get('status'))}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Tipo: <b>{texto_filtro(filtros.get('tipo'))}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Categoria: <b>{texto_filtro(filtros.get('categoria'))}</b>",
            subtitulo,
        )
    )
    elementos.append(
        Paragraph(
            f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} - "
            f"{len(pagamentos)} pagamento(s) encontrado(s)",
            subtitulo,
        )
    )
    elementos.append(Spacer(1, 8))

    tabela_resumo = Table(
        [
            ["Previsto", "Pago", "Pendente", "Pagos", "Pendentes", "Atrasados", "Vencem hoje"],
            [
                resumo.get("total_previsto_fmt", "R$ 0,00"),
                resumo.get("total_pago_fmt", "R$ 0,00"),
                resumo.get("total_pendente_fmt", "R$ 0,00"),
                str(resumo.get("qtd_pagos", 0)),
                str(resumo.get("qtd_pendentes", 0)),
                str(resumo.get("qtd_atrasados", 0)),
                str(resumo.get("qtd_vence_hoje", 0)),
            ],
        ],
        repeatRows=1,
        colWidths=[3.7 * cm, 3.7 * cm, 3.7 * cm, 2.6 * cm, 2.8 * cm, 2.8 * cm, 3.0 * cm],
    )
    tabela_resumo.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.75, colors.HexColor("#94a3b8")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#f8fafc")),
                ("FONTSIZE", (0, 0), (-1, -1), 8.4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    elementos.append(tabela_resumo)
    elementos.append(Spacer(1, 12))

    dados = [[
        "Grupo",
        "Vencimento",
        "Conta",
        "Categoria",
        "Subcategoria",
        "Previsto",
        "Pago",
        "Status",
    ]]

    if pagamentos:
        for item in pagamentos:
            estilo_linha = celula_pago if status_visual(item) == "PAGO" else celula
            dados.append(
                [
                    Paragraph(str(item.get("GRUPO_VENCIMENTO") or ""), estilo_linha),
                    Paragraph(str(item.get("DATA_VENCIMENTO_FMT") or ""), estilo_linha),
                    Paragraph(str(item.get("DESCRICAO") or ""), estilo_linha),
                    Paragraph(str(item.get("CATEGORIA") or ""), estilo_linha),
                    Paragraph(str(item.get("SUBCATEGORIA") or ""), estilo_linha),
                    Paragraph(str(item.get("VALOR_PREVISTO_FMT") or "R$ 0,00"), estilo_linha),
                    Paragraph(str(item.get("VALOR_PAGO_FMT") or "R$ 0,00"), estilo_linha),
                    Paragraph(str(item.get("STATUS_CALCULADO") or item.get("STATUS") or ""), estilo_linha),
                ]
            )
    else:
        dados.append(["", "", "Nenhum pagamento encontrado", "", "", "", "", ""])

    tabela = Table(
        dados,
        repeatRows=1,
        colWidths=[2.5 * cm, 2.2 * cm, 6.2 * cm, 3.3 * cm, 3.8 * cm, 2.7 * cm, 2.7 * cm, 3.2 * cm],
    )

    estilos_tabela = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.55, colors.HexColor("#94a3b8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
        ("BOX", (0, 0), (-1, -1), 0.85, colors.HexColor("#64748b")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("FONTSIZE", (0, 0), (-1, -1), 7.4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]

    for indice, item in enumerate(pagamentos, start=1):
        status_item = status_visual(item)
        if status_item == "PAGO":
            estilos_tabela.extend(
                [
                    ("BACKGROUND", (0, indice), (-1, indice), colors.HexColor("#dcfce7")),
                    ("TEXTCOLOR", (0, indice), (-1, indice), colors.HexColor("#065f46")),
                    ("LINEBEFORE", (0, indice), (0, indice), 3, colors.HexColor("#16a34a")),
                ]
            )
        elif status_item == "ATRASADO":
            estilos_tabela.extend(
                [
                    ("BACKGROUND", (0, indice), (-1, indice), colors.HexColor("#fef2f2")),
                    ("TEXTCOLOR", (0, indice), (-1, indice), colors.HexColor("#7f1d1d")),
                    ("LINEBEFORE", (0, indice), (0, indice), 3, colors.HexColor("#dc2626")),
                ]
            )
        elif status_item in {"VENCE HOJE", "PENDENTE"}:
            estilos_tabela.append(("LINEBEFORE", (0, indice), (0, indice), 3, colors.HexColor("#f59e0b")))

    tabela.setStyle(TableStyle(estilos_tabela))
    elementos.append(tabela)

    doc.build(elementos)
    buffer.seek(0)

    filename = nome_arquivo_relatorio(filtros, "pdf")
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/pagamentos/relatorio/xlsx")
async def relatorio_pagamentos_xlsx(
    ano: Optional[int] = Query(None),
    mes: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    tipo: Optional[str] = Query(None),
    categoria: Optional[str] = Query(None),
):
    filtros = montar_filtros_relatorio(
        ano=ano,
        mes=mes,
        status=status,
        tipo=tipo,
        categoria=categoria,
    )

    pagamentos = listar_pagamentos(filtros)
    resumo = montar_resumo_pagamentos(pagamentos)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        raise RuntimeError(
            "Para gerar Excel, instale a biblioteca openpyxl: pip install openpyxl"
        ) from exc

    def status_visual(item: dict) -> str:
        return str(item.get("STATUS_CALCULADO") or item.get("STATUS") or "").strip().upper()

    wb = Workbook()
    ws = wb.active
    ws.title = "Pagamentos"

    azul = "1D4ED8"
    azul_escuro = "0F172A"
    cinza = "F8FAFC"
    branco = "FFFFFF"
    verde_claro = "DCFCE7"
    verde_texto = "065F46"
    vermelho_claro = "FEE2E2"
    vermelho_texto = "7F1D1D"
    amarelo_claro = "FEF3C7"
    amarelo_texto = "92400E"

    borda_fina = Side(style="thin", color="94A3B8")
    borda_media = Side(style="medium", color="64748B")
    borda_verde = Side(style="medium", color="16A34A")
    borda_vermelha = Side(style="medium", color="DC2626")
    borda_amarela = Side(style="medium", color="F59E0B")

    def aplicar_borda(celula, side=borda_fina):
        celula.border = Border(top=side, bottom=side, left=side, right=side)

    ws.merge_cells("A1:H1")
    ws["A1"] = "Relatório de Pagamentos"
    ws["A1"].font = Font(bold=True, size=17, color=branco)
    ws["A1"].fill = PatternFill("solid", fgColor=azul_escuro)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:H2")
    ws["A2"] = (
        f"Ano: {texto_filtro(filtros.get('ano'))} | "
        f"Mês: {texto_filtro(filtros.get('mes'))} | "
        f"Status: {texto_filtro(filtros.get('status'))} | "
        f"Tipo: {texto_filtro(filtros.get('tipo'))} | "
        f"Categoria: {texto_filtro(filtros.get('categoria'))}"
    )
    ws["A2"].font = Font(size=10, color="475569", italic=True)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    ws.merge_cells("A3:H3")
    ws["A3"] = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} - {len(pagamentos)} pagamento(s) encontrado(s)"
    ws["A3"].font = Font(size=9, color="64748B")
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    resumo_cabecalhos = ["Previsto", "Pago", "Pendente", "Pagos", "Pendentes", "Atrasados", "Vencem hoje"]
    resumo_valores = [
        resumo.get("total_previsto_fmt", "R$ 0,00"),
        resumo.get("total_pago_fmt", "R$ 0,00"),
        resumo.get("total_pendente_fmt", "R$ 0,00"),
        resumo.get("qtd_pagos", 0),
        resumo.get("qtd_pendentes", 0),
        resumo.get("qtd_atrasados", 0),
        resumo.get("qtd_vence_hoje", 0),
    ]

    linha_resumo_header = 5
    for col, cab in enumerate(resumo_cabecalhos, start=1):
        cel = ws.cell(row=linha_resumo_header, column=col, value=cab)
        cel.font = Font(bold=True, color=branco)
        cel.fill = PatternFill("solid", fgColor=azul_escuro)
        cel.alignment = Alignment(horizontal="center", vertical="center")
        aplicar_borda(cel, borda_media)

    linha_resumo_valor = 6
    for col, valor in enumerate(resumo_valores, start=1):
        cel = ws.cell(row=linha_resumo_valor, column=col, value=valor)
        cel.font = Font(bold=True, color="111827")
        cel.fill = PatternFill("solid", fgColor=cinza)
        cel.alignment = Alignment(horizontal="center", vertical="center")
        aplicar_borda(cel, borda_fina)

    linha = 8
    cabecalhos = [
        "Grupo",
        "Vencimento",
        "Conta",
        "Categoria",
        "Subcategoria",
        "Previsto",
        "Pago",
        "Status",
    ]

    for col, cab in enumerate(cabecalhos, start=1):
        celula = ws.cell(row=linha, column=col, value=cab)
        celula.font = Font(bold=True, color=branco)
        celula.fill = PatternFill("solid", fgColor=azul)
        celula.alignment = Alignment(horizontal="center", vertical="center")
        aplicar_borda(celula, borda_media)
    ws.row_dimensions[linha].height = 24

    linha += 1

    if pagamentos:
        for item in pagamentos:
            valores = [
                item.get("GRUPO_VENCIMENTO", ""),
                item.get("DATA_VENCIMENTO_FMT", ""),
                item.get("DESCRICAO", ""),
                item.get("CATEGORIA", ""),
                item.get("SUBCATEGORIA", ""),
                item.get("VALOR_PREVISTO_FMT", "R$ 0,00"),
                item.get("VALOR_PAGO_FMT", "R$ 0,00"),
                item.get("STATUS_CALCULADO") or item.get("STATUS", ""),
            ]

            status_item = status_visual(item)
            if status_item == "PAGO":
                fill = PatternFill("solid", fgColor=verde_claro)
                font = Font(color=verde_texto, bold=True)
                left_side = borda_verde
            elif status_item == "ATRASADO":
                fill = PatternFill("solid", fgColor=vermelho_claro)
                font = Font(color=vermelho_texto, bold=True)
                left_side = borda_vermelha
            elif status_item in {"VENCE HOJE", "PENDENTE"}:
                fill = PatternFill("solid", fgColor=amarelo_claro if status_item == "VENCE HOJE" else branco)
                font = Font(color=amarelo_texto if status_item == "VENCE HOJE" else "111827")
                left_side = borda_amarela
            else:
                fill = PatternFill("solid", fgColor=cinza if linha % 2 == 0 else branco)
                font = Font(color="111827")
                left_side = borda_fina

            for col, valor in enumerate(valores, start=1):
                celula = ws.cell(row=linha, column=col, value=valor)
                celula.fill = fill
                celula.font = font
                celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                celula.border = Border(
                    top=borda_fina,
                    bottom=borda_fina,
                    left=left_side if col == 1 else borda_fina,
                    right=borda_fina,
                )

            ws.row_dimensions[linha].height = 34
            linha += 1
    else:
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=len(cabecalhos))
        cel = ws.cell(row=linha, column=1, value="Nenhum pagamento encontrado para os filtros aplicados.")
        cel.alignment = Alignment(horizontal="center", vertical="center")
        cel.fill = PatternFill("solid", fgColor=cinza)
        aplicar_borda(cel, borda_fina)

    larguras = [18, 14, 40, 20, 24, 16, 16, 18]
    for idx, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = largura

    max_linha = max(linha - 1, 8)
    ws.auto_filter.ref = f"A8:H{max_linha}"
    ws.freeze_panes = "A9"

    for linha_planilha in ws.iter_rows():
        for celula in linha_planilha:
            if celula.value is not None:
                celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = nome_arquivo_relatorio(filtros, "xlsx")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/pagamentos/salvar")
async def salvar_pagamento_rota(
    pagamento_id: str = Form(""),
    ano: str = Form(...),
    mes: str = Form(...),
    descricao: str = Form(...),
    tipo: str = Form(...),
    categoria: str = Form(...),
    subcategoria: str = Form(""),
    valor_previsto: str = Form(...),
    valor_pago: str = Form("0"),
    data_vencimento: str = Form(...),
    data_pagamento: str = Form(""),
    status: str = Form("PENDENTE"),
    recorrente: str = Form("NÃO"),
    conta_pagamento: str = Form(""),
    observacao: str = Form(""),
):
    try:
        salvar_pagamento(
            {
                "id": pagamento_id,
                "ano": ano,
                "mes": mes,
                "descricao": descricao,
                "tipo": tipo,
                "categoria": categoria,
                "subcategoria": subcategoria,
                "valor_previsto": valor_previsto,
                "valor_pago": valor_pago,
                "data_vencimento": data_vencimento,
                "data_pagamento": data_pagamento,
                "status": status,
                "recorrente": recorrente,
                "conta_pagamento": conta_pagamento,
                "observacao": observacao,
            }
        )
        return redirect_pagamentos(
            f"ano={ano}&mes={mes}&modo=cadastrar&mensagem=Pagamento salvo com sucesso"
        )
    except Exception as exc:
        return redirect_pagamentos(
            f"ano={ano}&mes={mes}&modo=cadastrar&erro={str(exc)}"
        )


@router.post("/pagamentos/{pagamento_id}/pagar")
async def marcar_pago_rota(pagamento_id: str):
    from core.financeiro.pagamentos_google import atualizar_status_pagamento

    try:
        atualizar_status_pagamento(pagamento_id, "PAGO")
        return redirect_pagamentos("mensagem=Pagamento marcado como pago e identificado como encontrado")
    except Exception as exc:
        return redirect_pagamentos(f"erro={str(exc)}")


@router.post("/pagamentos/{pagamento_id}/cancelar")
async def cancelar_pagamento_rota(pagamento_id: str):
    from core.financeiro.pagamentos_google import atualizar_status_pagamento

    try:
        atualizar_status_pagamento(pagamento_id, "CANCELADO")
        return redirect_pagamentos("mensagem=Pagamento cancelado")
    except Exception as exc:
        return redirect_pagamentos(f"erro={str(exc)}")


@router.post("/pagamentos/{pagamento_id}/excluir")
async def excluir_pagamento_rota(pagamento_id: str):
    try:
        excluir_pagamento(pagamento_id)
        return redirect_pagamentos("mensagem=Pagamento excluído")
    except Exception as exc:
        return redirect_pagamentos(f"erro={str(exc)}")


@router.post("/pagamentos/{pagamento_id}/duplicar")
async def duplicar_pagamento_rota(pagamento_id: str):
    try:
        novo = duplicar_pagamento(pagamento_id)
        ano = str(novo.get("ANO", "") or "")
        mes = str(novo.get("MES", "") or "")
        descricao = str(novo.get("DESCRICAO", "") or "")

        mensagem = f"Pagamento duplicado com sucesso para {mes}/{ano}: {descricao}"

        return redirect_pagamentos(f"ano={ano}&mes={mes}&mensagem={mensagem}")
    except Exception as exc:
        return redirect_pagamentos(f"erro={str(exc)}")



@router.post("/pagamentos/duplicar-mes")
async def duplicar_pagamentos_mes_rota(
    ano: str = Form(""),
    mes: str = Form(""),
):
    try:
        resultado = duplicar_pagamentos_mes(
            ano_origem=ano,
            mes_origem=mes,
        )

        ano_destino = str(resultado.get("ano_destino", ano) or ano)
        mes_destino = str(resultado.get("mes_destino", mes) or mes)
        mensagem = resultado.get("mensagem", "Pagamentos duplicados com sucesso.")

        return redirect_pagamentos(
            f"ano={ano_destino}&mes={mes_destino}&mensagem={mensagem}"
        )
    except Exception as exc:
        return redirect_pagamentos(
            f"ano={ano}&mes={mes}&erro={str(exc)}"
        )


@router.post("/pagamentos/conciliar")
async def conciliar_pagamentos_rota(
    ano: str = Form(""),
    mes: str = Form(""),
    status: str = Form(""),
    tipo: str = Form(""),
    categoria: str = Form(""),
):
    try:
        resultado = conciliar_pagamentos_com_lancamentos(
            {
                "ano": ano,
                "mes": mes,
                "status": status,
                "tipo": tipo,
                "categoria": categoria,
            },
            aplicar_alteracoes=True,
        )
        mensagem = resultado.get("mensagem", "Conferência concluída.")
        return redirect_pagamentos(f"ano={ano}&mes={mes}&mensagem={mensagem}")
    except Exception as exc:
        return redirect_pagamentos(f"ano={ano}&mes={mes}&modo=cadastrar&erro={str(exc)}")