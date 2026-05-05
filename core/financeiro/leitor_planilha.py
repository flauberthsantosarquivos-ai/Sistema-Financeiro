from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MESES_VALIDOS = {
    "JAN": "Janeiro",
    "FEV": "Fevereiro",
    "MAR": "Março",
    "ABR": "Abril",
    "MAI": "Maio",
    "JUN": "Junho",
    "JUL": "Julho",
    "AGO": "Agosto",
    "SET": "Setembro",
    "OUT": "Outubro",
    "NOV": "Novembro",
    "DEZ": "Dezembro",
}


ORDEM_MESES = [
    "JAN",
    "FEV",
    "MAR",
    "ABR",
    "MAI",
    "JUN",
    "JUL",
    "AGO",
    "SET",
    "OUT",
    "NOV",
    "DEZ",
]


ABAS_IGNORADAS = {
    "DADOS",
    "DADOS ",
    "HIPERCARD",
    "MODELO MENSAL",
    "RECEITA X DESPESA",
    "AMORTIZAÇÃO AP",
    "JAN 2025",
}


CATEGORIAS_PADRAO = {
    "PESSOAIS",
    "CASA",
    "ALIMENTAÇÃO",
    "SAÚDE",
    "EDUCAÇÃO",
    "TRANSPORTE",
    "COMUNICAÇÃO",
    "LAZER",
    "OUTROS",
}


def limpar_texto(valor: Any) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def para_float(valor: Any) -> float:
    if valor is None or valor == "":
        return 0.0

    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip()
    texto = texto.replace("R$", "").replace(" ", "")

    if texto.startswith("="):
        return 0.0

    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")

    try:
        return float(texto)
    except ValueError:
        return 0.0


def formatar_moeda(valor: float) -> str:
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def identificar_abas_mensais(wb) -> list[str]:
    abas = []

    for nome in wb.sheetnames:
        nome_limpo = nome.strip().upper()

        if nome_limpo in ABAS_IGNORADAS:
            continue

        if nome_limpo in MESES_VALIDOS:
            abas.append(nome)

    abas.sort(key=lambda nome: ORDEM_MESES.index(nome.strip().upper()))

    return abas


def filtrar_meses(
    abas_mensais: list[str],
    mes_inicio: str | None,
    mes_fim: str | None,
) -> list[str]:
    if not mes_inicio or not mes_fim:
        return abas_mensais

    mes_inicio = mes_inicio.strip().upper()
    mes_fim = mes_fim.strip().upper()

    if mes_inicio not in ORDEM_MESES or mes_fim not in ORDEM_MESES:
        return abas_mensais

    indice_inicio = ORDEM_MESES.index(mes_inicio)
    indice_fim = ORDEM_MESES.index(mes_fim)

    if indice_inicio > indice_fim:
        indice_inicio, indice_fim = indice_fim, indice_inicio

    meses_permitidos = set(ORDEM_MESES[indice_inicio : indice_fim + 1])

    return [
        aba
        for aba in abas_mensais
        if aba.strip().upper() in meses_permitidos
    ]


def ler_resumo_receita_x_despesa(wb) -> list[dict[str, Any]]:
    if "RECEITA X DESPESA" not in wb.sheetnames:
        return []

    ws = wb["RECEITA X DESPESA"]
    dados = []

    for row in range(2, ws.max_row + 1):
        mes = limpar_texto(ws.cell(row=row, column=1).value).upper()

        if not mes or mes == "TOTAL":
            continue

        if mes not in MESES_VALIDOS:
            continue

        receitas = para_float(ws.cell(row=row, column=2).value)
        despesas = para_float(ws.cell(row=row, column=3).value)
        saldo = receitas - despesas

        montante_inicio = para_float(ws.cell(row=row, column=5).value)
        montante_final = para_float(ws.cell(row=row, column=6).value)
        saldo_montante = montante_final - montante_inicio

        dados.append(
            {
                "mes": mes,
                "mes_nome": MESES_VALIDOS.get(mes, mes),
                "receitas": receitas,
                "despesas": despesas,
                "saldo": saldo,
                "montante_inicio": montante_inicio,
                "montante_final": montante_final,
                "saldo_montante": saldo_montante,
                "receitas_fmt": formatar_moeda(receitas),
                "despesas_fmt": formatar_moeda(despesas),
                "saldo_fmt": formatar_moeda(saldo),
                "montante_inicio_fmt": formatar_moeda(montante_inicio),
                "montante_final_fmt": formatar_moeda(montante_final),
                "saldo_montante_fmt": formatar_moeda(saldo_montante),
            }
        )

    dados.sort(key=lambda item: ORDEM_MESES.index(item["mes"]))

    return dados


def ler_resumo_categoria_da_aba(ws, mes: str) -> list[dict[str, Any]]:
    """
    Lê o resumo por categoria da própria planilha.

    Na planilha analisada:
    F5:G14

    F = categoria
    G = valor
    """
    categorias = []

    for row in range(5, 15):
        categoria = limpar_texto(ws.cell(row=row, column=6).value).upper()
        valor = para_float(ws.cell(row=row, column=7).value)

        if not categoria:
            continue

        if categoria in {"TOTAL", "TOTAL GERAL"}:
            continue

        if categoria not in CATEGORIAS_PADRAO:
            continue

        categorias.append(
            {
                "mes": mes,
                "mes_nome": MESES_VALIDOS.get(mes, mes),
                "categoria": categoria,
                "valor": valor,
                "valor_fmt": formatar_moeda(valor),
                "origem": "resumo",
            }
        )

    return categorias


def ler_lancamentos_detalhados_da_aba(ws, mes: str) -> list[dict[str, Any]]:
    """
    Lê o bloco DESPESAS DETALHADAS POR TIPO.

    Na planilha analisada:
    F18:I90

    F = descrição
    G = valor
    H = situação
    I = vencimento
    """
    lancamentos = []
    categoria_atual = ""

    for row in range(18, 91):
        descricao = limpar_texto(ws.cell(row=row, column=6).value)
        valor = para_float(ws.cell(row=row, column=7).value)
        situacao = limpar_texto(ws.cell(row=row, column=8).value)
        vencimento = limpar_texto(ws.cell(row=row, column=9).value)

        if not descricao:
            continue

        descricao_upper = descricao.upper()

        if descricao_upper in CATEGORIAS_PADRAO:
            categoria_atual = descricao_upper

            if valor == 0:
                continue

        if descricao_upper in {"TOTAL", "TOTAL GERAL"}:
            continue

        lancamentos.append(
            {
                "mes": mes,
                "mes_nome": MESES_VALIDOS.get(mes, mes),
                "categoria": categoria_atual or "SEM CATEGORIA",
                "descricao": descricao,
                "valor": valor,
                "valor_fmt": formatar_moeda(valor),
                "situacao": situacao,
                "vencimento": vencimento,
            }
        )

    return lancamentos


def recalcular_resumo_categoria_pelos_lancamentos(
    lancamentos: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    resumo = {}

    for item in lancamentos:
        mes = item["mes"]
        categoria = item["categoria"]

        if categoria not in CATEGORIAS_PADRAO:
            continue

        chave = (mes, categoria)

        if chave not in resumo:
            resumo[chave] = {
                "mes": mes,
                "mes_nome": MESES_VALIDOS.get(mes, mes),
                "categoria": categoria,
                "valor": 0.0,
                "origem": "detalhado",
            }

        resumo[chave]["valor"] += item["valor"]

    categorias = []

    for item in resumo.values():
        item["valor_fmt"] = formatar_moeda(item["valor"])
        categorias.append(item)

    categorias.sort(
        key=lambda item: (
            ORDEM_MESES.index(item["mes"]),
            item["categoria"],
        )
    )

    return categorias


def combinar_categorias_resumo_e_detalhado(
    categorias_resumo: list[dict[str, Any]],
    categorias_detalhado: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """
    Regra:
    - Usa o valor detalhado quando ele for maior que zero.
    - Se o detalhado estiver zerado ou ausente, usa o valor do resumo.
    - Isso resolve o caso de OUTROS estar preenchido no resumo, mas não aparecer nos lançamentos.
    """
    combinado = {}

    for item in categorias_resumo:
        chave = (item["mes"], item["categoria"])
        combinado[chave] = {
            "mes": item["mes"],
            "mes_nome": item["mes_nome"],
            "categoria": item["categoria"],
            "valor": item["valor"],
            "origem": "resumo",
        }

    for item in categorias_detalhado:
        chave = (item["mes"], item["categoria"])

        if chave not in combinado:
            combinado[chave] = {
                "mes": item["mes"],
                "mes_nome": item["mes_nome"],
                "categoria": item["categoria"],
                "valor": item["valor"],
                "origem": "detalhado",
            }
        else:
            if item["valor"] > 0:
                combinado[chave]["valor"] = item["valor"]
                combinado[chave]["origem"] = "detalhado"

    categorias = []

    for item in combinado.values():
        item["valor_fmt"] = formatar_moeda(item["valor"])
        categorias.append(item)

    categorias.sort(
        key=lambda item: (
            ORDEM_MESES.index(item["mes"]),
            item["categoria"],
        )
    )

    return categorias


def ler_patrimonio_da_aba(ws, mes: str) -> list[dict[str, Any]]:
    """
    Na planilha analisada, o demonstrativo financeiro fica em:
    K8:M20

    K = conta
    L = início do mês
    M = fim do mês
    """
    patrimonio = []

    for row in range(8, 21):
        conta = limpar_texto(ws.cell(row=row, column=11).value)
        inicio = para_float(ws.cell(row=row, column=12).value)
        fim = para_float(ws.cell(row=row, column=13).value)

        if not conta:
            continue

        patrimonio.append(
            {
                "mes": mes,
                "mes_nome": MESES_VALIDOS.get(mes, mes),
                "conta": conta,
                "inicio": inicio,
                "fim": fim,
                "inicio_fmt": formatar_moeda(inicio),
                "fim_fmt": formatar_moeda(fim),
            }
        )

    return patrimonio


def atualizar_resumo_mensal_com_categorias(
    resumo_mensal: list[dict[str, Any]],
    categorias: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    despesas_por_mes = {}

    for item in categorias:
        mes = item["mes"]
        despesas_por_mes[mes] = despesas_por_mes.get(mes, 0.0) + item["valor"]

    for item in resumo_mensal:
        mes = item["mes"]

        if mes in despesas_por_mes:
            item["despesas"] = despesas_por_mes[mes]
            item["saldo"] = item["receitas"] - item["despesas"]
            item["despesas_fmt"] = formatar_moeda(item["despesas"])
            item["saldo_fmt"] = formatar_moeda(item["saldo"])

    return resumo_mensal


def processar_planilha_financeira(
    caminho_arquivo: str | Path,
    mes_inicio: str | None = None,
    mes_fim: str | None = None,
) -> dict[str, Any]:
    caminho = Path(caminho_arquivo)

    if not caminho.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

    wb = load_workbook(caminho, data_only=True)

    abas_mensais = identificar_abas_mensais(wb)
    abas_mensais = filtrar_meses(
        abas_mensais=abas_mensais,
        mes_inicio=mes_inicio,
        mes_fim=mes_fim,
    )

    resumo_mensal = ler_resumo_receita_x_despesa(wb)

    if abas_mensais:
        meses_selecionados = {aba.strip().upper() for aba in abas_mensais}
        resumo_mensal = [
            item
            for item in resumo_mensal
            if item["mes"] in meses_selecionados
        ]

    categorias_resumo = []
    lancamentos = []
    patrimonio = []

    for aba in abas_mensais:
        ws = wb[aba]
        mes = aba.strip().upper()

        categorias_resumo.extend(ler_resumo_categoria_da_aba(ws, mes))
        lancamentos.extend(ler_lancamentos_detalhados_da_aba(ws, mes))
        patrimonio.extend(ler_patrimonio_da_aba(ws, mes))

    categorias_detalhado = recalcular_resumo_categoria_pelos_lancamentos(lancamentos)

    categorias = combinar_categorias_resumo_e_detalhado(
        categorias_resumo=categorias_resumo,
        categorias_detalhado=categorias_detalhado,
    )

    resumo_mensal = atualizar_resumo_mensal_com_categorias(
        resumo_mensal=resumo_mensal,
        categorias=categorias,
    )

    total_receitas = sum(item["receitas"] for item in resumo_mensal)
    total_despesas = sum(item["despesas"] for item in resumo_mensal)
    saldo_total = total_receitas - total_despesas

    despesas_por_categoria = {}

    for item in categorias:
        categoria = item["categoria"]
        despesas_por_categoria[categoria] = despesas_por_categoria.get(categoria, 0.0) + item["valor"]

    ranking_categorias = [
        {
            "categoria": categoria,
            "valor": valor,
            "valor_fmt": formatar_moeda(valor),
        }
        for categoria, valor in sorted(
            despesas_por_categoria.items(),
            key=lambda item: item[1],
            reverse=True,
        )
    ]

    maior_categoria = (
        ranking_categorias[0]
        if ranking_categorias
        else {
            "categoria": "-",
            "valor": 0.0,
            "valor_fmt": formatar_moeda(0.0),
        }
    )

    patrimonio_final = 0.0
    patrimonio_por_mes = {}

    for item in patrimonio:
        if item["conta"].upper() == "TOTAL":
            patrimonio_por_mes[item["mes"]] = item["fim"]

    if patrimonio_por_mes:
        ultimo_mes = sorted(
            patrimonio_por_mes.keys(),
            key=lambda mes: ORDEM_MESES.index(mes),
        )[-1]
        patrimonio_final = patrimonio_por_mes[ultimo_mes]

    percentual_comprometimento = 0.0

    if total_receitas > 0:
        percentual_comprometimento = (total_despesas / total_receitas) * 100

    periodo_descricao = "Todos os meses"

    if mes_inicio and mes_fim:
        mes_inicio = mes_inicio.upper()
        mes_fim = mes_fim.upper()

        if mes_inicio == mes_fim:
            periodo_descricao = MESES_VALIDOS.get(mes_inicio, mes_inicio)
        else:
            periodo_descricao = (
                f"{MESES_VALIDOS.get(mes_inicio, mes_inicio)} "
                f"a {MESES_VALIDOS.get(mes_fim, mes_fim)}"
            )

    return {
        "arquivo": caminho.name,
        "periodo_descricao": periodo_descricao,
        "abas_mensais": abas_mensais,
        "cards": {
            "total_receitas": total_receitas,
            "total_despesas": total_despesas,
            "saldo_total": saldo_total,
            "patrimonio_final": patrimonio_final,
            "percentual_comprometimento": percentual_comprometimento,
            "maior_categoria": maior_categoria,
            "total_receitas_fmt": formatar_moeda(total_receitas),
            "total_despesas_fmt": formatar_moeda(total_despesas),
            "saldo_total_fmt": formatar_moeda(saldo_total),
            "patrimonio_final_fmt": formatar_moeda(patrimonio_final),
            "percentual_comprometimento_fmt": f"{percentual_comprometimento:.2f}%",
        },
        "resumo_mensal": resumo_mensal,
        "categorias": categorias,
        "ranking_categorias": ranking_categorias,
        "lancamentos": lancamentos,
        "patrimonio": patrimonio,
    }